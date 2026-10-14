import os
import csv
import tempfile
from collections import defaultdict
from datetime import datetime, date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference


# ── Colour palette (openpyxl ARGB) ──────────────────────────────────────────
NAVY     = "FF0F2351"
BLUE     = "FF1D4ED8"
LTBLUE   = "FFDBEAFE"
TEAL     = "FF0F766E"
TEAL_BG  = "FFCCFBF1"
GREEN    = "FF166534"
GREEN_BG = "FFDCFCE7"
AMBER    = "FF92400E"
AMBER_BG = "FFFEF3C7"
RED      = "FF991B1B"
RED_BG   = "FFFEE2E2"
LGRAY    = "FFF3F4F6"
WHITE    = "FFFFFFFF"
DGRAY    = "FF374151"
MIDGRAY  = "FF6B7280"


# ── Style helpers ────────────────────────────────────────────────────────────
def _side(c="D1D5DB"):
    return Side(style="thin", color=c)

def _border(c="D1D5DB"):
    s = _side(c)
    return Border(left=s, right=s, top=s, bottom=s)

def _fill(argb):
    return PatternFill("solid", fgColor=argb)

def _font(bold=False, color="FF374151", size=11, italic=False):
    return Font(name="Arial", bold=bold, color=color, size=size, italic=italic)

def _align(h="left", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def _set_widths(ws, widths: dict):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

def _header_row(ws, row: int, headers: list, fill=NAVY):
    for ci, text in enumerate(headers, 1):
        c = ws.cell(row=row, column=ci, value=text)
        c.font      = _font(bold=True, color="FFFFFFFF", size=11)
        c.fill      = _fill(fill)
        c.alignment = _align("center")
        c.border    = _border("1D4ED8")
    ws.row_dimensions[row].height = 22

def _status_fill(s):
    return _fill({"Present": GREEN_BG, "Late": AMBER_BG, "Absent": RED_BG}.get(s, LGRAY))

def _status_font(s):
    return _font(bold=True, color={"Present": GREEN, "Late": AMBER, "Absent": RED}.get(s, DGRAY))

def _fmt_dt(v):
    """Flatten datetime / ISO-string / None → 'HH:MM:SS' or '—'."""
    if v is None:
        return "—"
    if isinstance(v, datetime):
        return v.strftime("%H:%M:%S")
    try:
        return datetime.fromisoformat(str(v)).strftime("%H:%M:%S")
    except Exception:
        return str(v) if v else "—"

def _fmt_conf(v):
    """Format confidence float or string → '87.3%'."""
    try:
        return f"{float(str(v).replace('%', '')):.1f}%"
    except (TypeError, ValueError):
        return "—"


# ════════════════════════════════════════════════════════════════════════════
class AttendanceManager:

    LATE_AFTER_HOUR   = 9
    LATE_AFTER_MINUTE = 15

    def __init__(self, db_session_factory):
        self.Session = db_session_factory

    def _models(self):
        from database.models import AttendanceRecord, Student
        return AttendanceRecord, Student

    # ── Consecutive Class Absences Rule ──────────────────────────────────────
    def apply_consecutive_absence_rules(self, session, date_str):
        """
        If a student is present in the 1st session of a consecutive pair,
        but not present in the 2nd session, they are marked as Absent for the 2nd session.
        """
        AttendanceRecord, _ = self._models()
        records = session.query(AttendanceRecord).filter_by(date=date_str).all()
        
        student_records = defaultdict(dict)
        for r in records:
            student_records[r.student_id][r.session] = r
                
        now = datetime.now()
        current_date = now.date().isoformat()
        current_hour = now.hour
        
        # Consecutive session pairs (Session 1, Session 2, and the hour after which Session 2 is over)
        consecutive_pairs = [
            ("Morning", "Afternoon", 16),
            ("Afternoon", "Evening", 20)
        ]
        
        modified = False
        for student_id, sessions in student_records.items():
            for s1, s2, over_hour in consecutive_pairs:
                # Is the session s2 over?
                is_over = False
                if date_str < current_date:
                    is_over = True
                elif date_str == current_date:
                    if current_hour >= over_hour:
                        is_over = True
                        
                if is_over:
                    rec1 = sessions.get(s1)
                    rec2 = sessions.get(s2)
                    
                    if rec1 and rec1.status in ["Present", "Late"]:
                        if not rec2:
                            # Create an Absent record for the 2nd session
                            new_rec = AttendanceRecord(
                                student_id=student_id,
                                student_name=rec1.student_name,
                                date=date_str,
                                check_in_time=None,
                                confidence=0.0,
                                status="Absent",
                                session=s2
                            )
                            session.add(new_rec)
                            sessions[s2] = new_rec
                            modified = True
                        elif rec2.status not in ["Present", "Late", "Absent"]:
                            rec2.status = "Absent"
                            modified = True
                            
        if modified:
            session.commit()

    # ── Mark Attendance ──────────────────────────────────────────────────────
    def mark_attendance(self, student_id, student_name, confidence):
        AttendanceRecord, _ = self._models()
        session = self.Session()
        try:
            today = date.today().isoformat()
            now   = datetime.now()
            
            # Map current hour to a session
            session_name = self._get_session_name(now)
                
            existing = (session.query(AttendanceRecord)
                        .filter_by(student_id=student_id, date=today, session=session_name).first())
            if existing is None:
                rec = AttendanceRecord(
                    student_id=student_id, student_name=student_name,
                    date=today, check_in_time=now, confidence=confidence,
                    status=self._compute_status(now, session_name),
                    session=session_name
                )
                session.add(rec)
                session.commit()
                self.apply_consecutive_absence_rules(session, today)
                return {"action": "check_in", "record": rec.to_dict()}
            else:
                if confidence > (existing.confidence or 0.0):
                    existing.confidence = confidence
                    session.commit()
                return {"action": "check_in", "record": existing.to_dict()}
        except Exception:
            session.rollback()
            raise
        finally:
            session.close()

    # ── Queries ──────────────────────────────────────────────────────────────
    def get_today_attendance(self):
        AttendanceRecord, _ = self._models()
        session = self.Session()
        try:
            today = date.today().isoformat()
            self.apply_consecutive_absence_rules(session, today)
            recs = session.query(AttendanceRecord).filter_by(date=today).all()
            return [r.to_dict() for r in recs]
        finally:
            session.close()

    def get_attendance_by_date(self, target_date):
        AttendanceRecord, _ = self._models()
        session = self.Session()
        try:
            self.apply_consecutive_absence_rules(session, target_date)
            recs = session.query(AttendanceRecord).filter_by(date=target_date).all()
            return [r.to_dict() for r in recs]
        finally:
            session.close()

    def get_student_history(self, student_id, limit=30):
        AttendanceRecord, _ = self._models()
        session = self.Session()
        try:
            recs = (session.query(AttendanceRecord)
                    .filter_by(student_id=student_id)
                    .order_by(AttendanceRecord.date.desc())
                    .limit(limit).all())
            return [r.to_dict() for r in recs]
        finally:
            session.close()

    def get_all_attendance(self):
        AttendanceRecord, _ = self._models()
        session = self.Session()
        try:
            recs = (session.query(AttendanceRecord)
                    .order_by(AttendanceRecord.date.desc()).all())
            return [r.to_dict() for r in recs]
        finally:
            session.close()

    # ── Stats ────────────────────────────────────────────────────────────────
    def get_stats(self):
        AttendanceRecord, Student = self._models()
        session = self.Session()
        try:
            today   = date.today().isoformat()
            self.apply_consecutive_absence_rules(session, today)
            total   = session.query(Student).filter_by(is_active=True).count()
            recs    = session.query(AttendanceRecord).filter_by(date=today).all()
            
            # Count unique students present or late
            present_ids = set(r.student_id for r in recs if r.status in ["Present", "Late"])
            present = len(present_ids)
            late    = len(set(r.student_id for r in recs if r.status == "Late"))
            absent  = max(0, total - present)
            rate    = round(present / total * 100, 1) if total else 0.0
            return {
                "total_students": total, "present_today": present,
                "absent_today": absent, "late_today": late,
                "attendance_rate": rate,
            }
        finally:
            session.close()

    def get_weekly_trend(self):
        AttendanceRecord, _ = self._models()
        session = self.Session()
        try:
            today = date.today()
            result = []
            for i in range(6, -1, -1):
                d     = (today - timedelta(days=i)).isoformat()
                self.apply_consecutive_absence_rules(session, d)
                recs  = session.query(AttendanceRecord).filter_by(date=d).all()
                count = len(set(r.student_id for r in recs if r.status in ["Present", "Late"]))
                result.append({"date": d, "count": count})
            return result
        finally:
            session.close()

    # ── CSV Export ───────────────────────────────────────────────────────────
    def export_csv(self, target_date=None):
        records = (self.get_attendance_by_date(target_date)
                   if target_date else self.get_all_attendance())

        if not records:
            return None

        # FIX: Use system temp directory to prevent cloud permission 500 errors
        export_dir = os.path.join(tempfile.gettempdir(), "attendance_exports")
        os.makedirs(export_dir, exist_ok=True)
        label = target_date or "all"
        path  = os.path.join(export_dir, f"attendance_{label}.csv")

        headers = [
            "Student ID", "Student Name", "Date",
            "Check-In Time",
            "Confidence (%)", "Status", "Session",
        ]

        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            for r in records:
                writer.writerow([
                    r.get("student_id",   ""),
                    r.get("student_name", ""),
                    r.get("date",         ""),
                    _fmt_dt(r.get("check_in_time")),
                    _fmt_conf(r.get("confidence")),
                    r.get("status",  ""),
                    r.get("session", ""),
                ])
        return path

    # ── Excel Export ─────────────────────────────────────────────────────────
    def export_excel(self, target_date=None):
        records = (self.get_attendance_by_date(target_date)
                   if target_date else self.get_all_attendance())

        if not records:
            return None

        # FIX: Use system temp directory to prevent cloud permission 500 errors
        export_dir = os.path.join(tempfile.gettempdir(), "attendance_exports")
        os.makedirs(export_dir, exist_ok=True)
        label = target_date or "all"
        path  = os.path.join(export_dir, f"attendance_{label}.xlsx")

        wb = Workbook()
        wb.remove(wb.active)

        self._sheet_summary(wb, target_date)
        self._sheet_daily(wb, records)
        self._sheet_breakdown(wb, records)

        wb.save(path)
        return path

    # ── Sheet 1: Summary ─────────────────────────────────────────────────────
    def _sheet_summary(self, wb, target_date):
        ws = wb.create_sheet("Summary")
        ws.sheet_view.showGridLines = False

        # Title
        ws.merge_cells("A1:F1")
        c = ws["A1"]
        c.value     = "Smart Face Recognition — Attendance Report"
        c.font      = _font(bold=True, color="FFFFFFFF", size=15)
        c.fill      = _fill(NAVY)
        c.alignment = _align("center")
        ws.row_dimensions[1].height = 36

        ws.merge_cells("A2:F2")
        sub = ws["A2"]
        sub.value     = (f"Date: {target_date}" if target_date
                         else f"All Records  |  Generated: {date.today()}")
        sub.font      = _font(italic=True, color=MIDGRAY, size=10)
        sub.fill      = _fill(LGRAY)
        sub.alignment = _align("center")
        ws.row_dimensions[2].height = 18

        # Stat cards row 4-5
        stats = self.get_stats()
        cards = [
            ("Total Students",  stats["total_students"],          BLUE,  LTBLUE),
            ("Present Today",   stats["present_today"],           GREEN, GREEN_BG),
            ("Absent Today",    stats["absent_today"],            RED,   RED_BG),
            ("Late Today",      stats["late_today"],              AMBER, AMBER_BG),
            ("Attendance Rate", f"{stats['attendance_rate']}%",    TEAL,  TEAL_BG),
        ]
        ws.row_dimensions[3].height = 6
        for i, (label, value, tc, bg) in enumerate(cards):
            col = get_column_letter(i + 1)
            lc = ws[f"{col}4"]
            lc.value = label; lc.font = _font(bold=True, color=tc, size=10)
            lc.fill = _fill(bg); lc.alignment = _align("center"); lc.border = _border()
            ws.row_dimensions[4].height = 20
            vc = ws[f"{col}5"]
            vc.value = value
            vc.font = Font(name="Arial", bold=True, color=tc, size=20)
            vc.fill = _fill(bg); vc.alignment = _align("center"); vc.border = _border()
            ws.row_dimensions[5].height = 34
            ws.column_dimensions[col].width = 18

        # Weekly trend table (A8 onward — source for chart)
        ws["A7"] = "Date"; ws["B7"] = "Students Present"
        ws["A7"].font = ws["B7"].font = _font(bold=True, color="FFFFFFFF")
        ws["A7"].fill = ws["B7"].fill = _fill(NAVY)
        ws["A7"].alignment = ws["B7"].alignment = _align("center")
        trend = self.get_weekly_trend()
        for i, item in enumerate(trend, start=8):
            ws[f"A{i}"] = item["date"]
            ws[f"B{i}"] = item["count"]

        # Bar chart
        chart = BarChart()
        chart.type = "col"; chart.style = 10
        chart.title = "7-Day Attendance Trend"
        chart.y_axis.title = "Students Present"
        chart.x_axis.title = "Date"
        chart.height = 12; chart.width = 22
        data = Reference(ws, min_col=2, min_row=7, max_row=7 + len(trend))
        cats = Reference(ws, min_col=1, min_row=8, max_row=7 + len(trend))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.series[0].graphicalProperties.solidFill = "1D4ED8"
        ws.add_chart(chart, "A10")

    # ── Sheet 2: Daily Attendance ────────────────────────────────────────────
    def _sheet_daily(self, wb, records):
        ws = wb.create_sheet("Daily Attendance")
        ws.sheet_view.showGridLines = False

        ws.merge_cells("A1:H1")
        t = ws["A1"]
        t.value = "Daily Attendance Log"
        t.font = _font(bold=True, color="FFFFFFFF", size=13)
        t.fill = _fill(NAVY); t.alignment = _align("center")
        ws.row_dimensions[1].height = 30

        cols = ["Student ID","Student Name","Date","Check-In",
                "Confidence (%)","Status","Session"]
        _header_row(ws, 2, cols, BLUE)
        ws.freeze_panes = "A3"
        ws.auto_filter.ref = f"A2:G{len(records)+2}"

        for ri, rec in enumerate(records, start=3):
            status = rec.get("status", "")
            sf = _status_fill(status)
            row = [
                rec.get("student_id", ""),
                rec.get("student_name", ""),
                rec.get("date", ""),
                _fmt_dt(rec.get("check_in_time")),
                _fmt_conf(rec.get("confidence")),
                status,
                rec.get("session", ""),
            ]
            for ci, val in enumerate(row, 1):
                c = ws.cell(row=ri, column=ci, value=val)
                c.fill = sf; c.border = _border()
                c.alignment = _align("center" if ci != 2 else "left")
                c.font = _status_font(status) if ci == 6 else _font(size=10)
            ws.row_dimensions[ri].height = 18

        _set_widths(ws, {"A":14,"B":24,"C":14,"D":14,"E":16,"F":12,"G":14})

        # Totals row
        tr = len(records) + 4
        ws.merge_cells(f"A{tr}:C{tr}")
        sc = ws[f"A{tr}"]
        sc.value = f"Total Records: {len(records)}"
        sc.font = _font(bold=True, color=BLUE, size=11)

        p = sum(1 for r in records if r.get("status") == "Present")
        l = sum(1 for r in records if r.get("status") == "Late")
        a = sum(1 for r in records if r.get("status") == "Absent")
        ws[f"F{tr}"].value = f"P:{p}  L:{l}  A:{a}"
        ws[f"F{tr}"].font  = _font(bold=True, size=10)

    # ── Sheet 3: Student Breakdown ───────────────────────────────────────────
    def _sheet_breakdown(self, wb, records):
        ws = wb.create_sheet("Student Breakdown")
        ws.sheet_view.showGridLines = False

        ws.merge_cells("A1:F1")
        t = ws["A1"]
        t.value = "Per-Student Attendance Summary"
        t.font = _font(bold=True, color="FFFFFFFF", size=13)
        t.fill = _fill(TEAL); t.alignment = _align("center")
        ws.row_dimensions[1].height = 30

        _header_row(ws, 2, ["Student ID","Student Name","Total Days",
                             "Present","Late","Avg Confidence"], TEAL)
        ws.freeze_panes = "A3"

        # Aggregate
        agg = defaultdict(lambda: {
            "name":"", "total":0, "present":0, "late":0, "conf":[], })
        for rec in records:
            sid = rec.get("student_id","")
            agg[sid]["name"]  = rec.get("student_name","")
            agg[sid]["total"] += 1
            s = rec.get("status","")
            if s == "Present": agg[sid]["present"] += 1
            elif s == "Late":  agg[sid]["late"]    += 1
            try:
                agg[sid]["conf"].append(
                    float(str(rec.get("confidence","0")).replace("%","")))
            except ValueError:
                pass

        # FIX: Added `or ""` to prevent 500 error TypeErrors if a student_name is None
        for ri, (sid, d) in enumerate(
                sorted(agg.items(), key=lambda x: x[1]["name"] or ""), start=3):
            avg = sum(d["conf"]) / len(d["conf"]) if d["conf"] else 0.0
            rf = _fill(LGRAY if ri % 2 == 0 else WHITE)
            row = [sid, d["name"], d["total"], d["present"],
                   d["late"], f"{avg:.1f}%"]
            for ci, val in enumerate(row, 1):
                c = ws.cell(row=ri, column=ci, value=val)
                c.fill = rf; c.border = _border(); c.font = _font(size=10)
                c.alignment = _align("center" if ci != 2 else "left")
            ws.row_dimensions[ri].height = 18

        _set_widths(ws, {"A":14,"B":24,"C":13,"D":12,"E":10,"F":18})

    # ── Private helpers ──────────────────────────────────────────────────────
    def _compute_status(self, dt, session_name):
        if session_name == "Morning":
            # Late after 09:15
            late_hour, late_min = 9, 15
        elif session_name == "Afternoon":
            # Late after 13:15 (assuming 13:00 start)
            late_hour, late_min = 13, 15
        elif session_name == "Evening":
            # Late after 17:15 (assuming 17:00 start)
            late_hour, late_min = 17, 15
        else:
            return "Present"
            
        if (dt.hour > late_hour or
                (dt.hour == late_hour
                 and dt.minute >= late_min)):
            return "Late"
        return "Present"

    def _get_session_name(self, dt):
        if dt.hour < 12:
            return "Morning"
        elif dt.hour < 16:
            return "Afternoon"
        else:
            return "Evening"
    # At the very bottom of utils/attendance_manager.py

if __name__ == "__main__":
    # 1. Create a dummy session factory for testing
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    
    # Replace with your actual test database URL
    engine = create_engine('sqlite:///../test_attendance.db') 
    TestSession = sessionmaker(bind=engine)
    
    # 2. Instantiate your manager
    manager = AttendanceManager(TestSession)
    
    # 3. Test a specific function
    print("Testing Excel Export...")
    try:
        file_path = manager.export_excel()
        print(f"Success! Excel file generated at: {file_path}")
    except Exception as e:
        print(f"An error occurred: {e}")